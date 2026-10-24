"""
Script 03: Build the Excel reporting workbook.

Purpose
    Build workbook/K12_Student_Analytics_Report.xlsx from the processed
    tables. The KPI Summary sheet calculates its values with Excel
    formulas over the Student Summary sheet, so the workbook stays live
    if the data changes. The two group summary sheets hold values
    calculated by this script and are labeled as pivot ready summary
    tables. They are not PivotTables. The original portfolio workbook,
    which does contain PivotTables, is preserved in workbook/original.

Sheets
    Read Me, KPI Summary, Student Summary, Attendance By Grade,
    Performance By Subject, Data Dictionary, Quality Summary.

Run order
    Run after 02_clean_and_validate.py.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE_DIR = Path(__file__).resolve().parent.parent
PROCESSED_DIR = BASE_DIR / "data" / "processed"
WORKBOOK_PATH = BASE_DIR / "workbook" / "K12_Student_Analytics_Report.xlsx"

NAVY_FILL = PatternFill("solid", fgColor="16324F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13, color="16324F")


def style_header_row(sheet, row, last_column):
    for column in range(1, last_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = NAVY_FILL
        cell.font = HEADER_FONT


def set_widths(sheet, widths):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_frame(sheet, frame, start_row=1):
    for column_index, column_name in enumerate(frame.columns, start=1):
        sheet.cell(row=start_row, column=column_index, value=column_name)
    style_header_row(sheet, start_row, len(frame.columns))
    for row_index, (_, row) in enumerate(frame.iterrows(), start=start_row + 1):
        for column_index, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            sheet.cell(row=row_index, column=column_index, value=value)
    sheet.freeze_panes = sheet.cell(row=start_row + 1, column=1)


def build_workbook():
    summary = pd.read_csv(PROCESSED_DIR / "fact_student_summary.csv")
    attendance = pd.read_csv(PROCESSED_DIR / "fact_attendance.csv")
    assessments = pd.read_csv(PROCESSED_DIR / "fact_assessment.csv")
    validation = pd.read_csv(BASE_DIR / "outputs" / "validation_results.csv")
    for column in ["Enrollment_Date", "Assessment_Date"]:
        summary[column] = pd.to_datetime(summary[column]).dt.date

    workbook = Workbook()

    # ------------------------------------------------------------- Read Me
    read_me = workbook.active
    read_me.title = "Read Me"
    read_me.column_dimensions["A"].width = 110
    lines = [
        ("K12 Student Analytics Report", True),
        ("", False),
        ("Purpose", True),
        ("This workbook reports the headline results of the K12 student data analytics project in Excel. It is built by scripts/03_build_workbook.py from the processed tables in data/processed, so it can be rebuilt after any data refresh.", False),
        ("", False),
        ("Analysis population", True),
        ("Every sheet describes the same population: the 110 students listed as Active in the student table. Seven Withdrawn students are excluded because their records cover a part year. Two Active students have no attendance row and three have no assessment row, so some averages rest on 108 or 107 students. The KPI Summary notes state the base for each figure.", False),
        ("", False),
        ("How the sheets work", True),
        ("KPI Summary holds live Excel formulas over the Student Summary sheet. Change a value in Student Summary and the KPIs update. Attendance By Grade and Performance By Subject hold values calculated by the build script. They are pivot ready summary tables, ready to chart or to rebuild as PivotTables. They are not PivotTables themselves.", False),
        ("The original portfolio workbook, which contains PivotTables built by hand in Excel, is preserved unchanged in workbook/original.", False),
        ("", False),
        ("Sources", True),
        ("The source extracts sit in data/source. The data is simulated practice data for one school, grades 5 to 8, in the 2025 to 2026 school year. It is not data about any real school or any real student.", False),
        ("", False),
        ("Definitions that matter", True),
        ("Met means a score of 70 or above. Below means a score under 70. The chronic absence group is a project definition: an attendance rate below 90 percent. The attendance bands are project categories: below 85 is At Risk, 85 to 89.9 is Needs Monitoring, 90 to 94.9 is On Track, and 95 and above is Strong.", False),
    ]
    for row_index, (text, is_title) in enumerate(lines, start=1):
        cell = read_me.cell(row=row_index, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if is_title:
            cell.font = TITLE_FONT

    # ---------------------------------------------------------- KPI Summary
    kpi_sheet = workbook.create_sheet("KPI Summary")
    last_row = len(summary) + 1
    s = f"'Student Summary'"
    kpi_rows = [
        ("Total Students", f"=COUNTA({s}!$A$2:$A${last_row})", "0",
         "Students in the analysis population."),
        ("Students With Attendance", f"=COUNT({s}!$H$2:$H${last_row})", "0",
         "Students with an attendance rate recorded."),
        ("Average Attendance Rate", f"=AVERAGE({s}!$H$2:$H${last_row})", "0.00",
         "Average of individual student rates. Each student counts equally."),
        ("Students With Assessment", f"=COUNT({s}!$M$2:$M${last_row})", "0",
         "Students with a score recorded."),
        ("Average Score", f"=AVERAGE({s}!$M$2:$M${last_row})", "0.00",
         "Average score across assessed students."),
        ("Students Met Benchmark", f"=COUNTIF({s}!$N$2:$N${last_row},\"Met\")", "0",
         "Assessed students with a score of 70 or above."),
        ("Students Below Benchmark", f"=COUNTIF({s}!$N$2:$N${last_row},\"Below\")", "0",
         "Assessed students with a score under 70."),
        ("Percent Met Benchmark", "=B7/B5", "0.0%",
         "Students Met Benchmark divided by Students With Assessment."),
        ("Chronic Absence Count", f"=COUNTIF({s}!$J$2:$J${last_row},\"Yes\")", "0",
         "Students with an attendance rate below 90 percent, the project definition."),
        ("Chronic Absence Rate", "=B10/B3", "0.0%",
         "Chronic Absence Count divided by Students With Attendance."),
        ("Total Incidents", f"=SUM({s}!$O$2:$O${last_row})", "0",
         "Total discipline incidents across students with a discipline row."),
        ("Students With Incidents", f"=COUNTIF({s}!$O$2:$O${last_row},\">0\")", "0",
         "Students with at least one recorded incident."),
        ("Average Incidents Per Student",
         f"=B12/COUNT({s}!$O$2:$O${last_row})", "0.00",
         "Total Incidents divided by students with a discipline row."),
    ]
    kpi_sheet.cell(row=1, column=1, value="KPI")
    kpi_sheet.cell(row=1, column=2, value="Value")
    kpi_sheet.cell(row=1, column=3, value="Definition and base")
    style_header_row(kpi_sheet, 1, 3)
    for index, (name, formula, number_format, note) in enumerate(kpi_rows, start=2):
        kpi_sheet.cell(row=index, column=1, value=name)
        value_cell = kpi_sheet.cell(row=index, column=2, value=formula)
        value_cell.number_format = number_format
        note_cell = kpi_sheet.cell(row=index, column=3, value=note)
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")
    kpi_sheet.freeze_panes = "A2"
    set_widths(kpi_sheet, [30, 12, 78])

    # -------------------------------------------------------- Student Summary
    summary_sheet = workbook.create_sheet("Student Summary")
    write_frame(summary_sheet, summary)
    table = Table(displayName="StudentSummary", ref=f"A1:P{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showRowStripes=True
    )
    summary_sheet.add_table(table)
    set_widths(summary_sheet, [12, 7, 13, 15, 13, 12, 12, 15, 17, 17, 11, 15, 8, 11, 10, 15])
    for row_index in range(2, last_row + 1):
        for column in ["D", "L"]:
            summary_sheet[f"{column}{row_index}"].number_format = "yyyy-mm-dd"

    # ---------------------------------------------------- Attendance By Grade
    by_grade = (
        attendance.groupby("Grade")
        .agg(
            Students=("Student_ID", "count"),
            Average_Attendance_Rate=("Attendance_Rate", "mean"),
            Average_Days_Absent=("Days_Absent", "mean"),
        )
        .reset_index()
    )
    chronic = (
        attendance.groupby("Grade")["Is_Chronic_Absence"]
        .apply(lambda values: round(100 * (values == "Yes").mean(), 1))
        .rename("Chronic_Absence_Rate_Percent")
        .reset_index()
    )
    by_grade = by_grade.merge(chronic, on="Grade")
    by_grade["Average_Attendance_Rate"] = by_grade["Average_Attendance_Rate"].round(2)
    by_grade["Average_Days_Absent"] = by_grade["Average_Days_Absent"].round(1)
    grade_sheet = workbook.create_sheet("Attendance By Grade")
    grade_sheet.cell(
        row=1, column=1,
        value="Pivot ready summary table. Values calculated by the build script from fact_attendance.",
    ).font = Font(italic=True, color="6B7280")
    write_frame(grade_sheet, by_grade, start_row=3)
    grade_sheet.freeze_panes = "A4"
    set_widths(grade_sheet, [8, 11, 24, 20, 26])

    # ------------------------------------------------ Performance By Subject
    by_subject = (
        assessments.groupby("Assessment")
        .agg(Students=("Student_ID", "count"), Average_Score=("Score", "mean"))
        .reset_index()
    )
    met = (
        assessments.groupby("Assessment")["Benchmark"]
        .apply(lambda values: round(100 * (values == "Met").mean(), 1))
        .rename("Percent_Met")
        .reset_index()
    )
    below = (
        assessments[assessments["Benchmark"] == "Below"]
        .groupby("Assessment")
        .size()
        .rename("Students_Below")
        .reset_index()
    )
    by_subject = by_subject.merge(met, on="Assessment").merge(below, on="Assessment")
    by_subject["Average_Score"] = by_subject["Average_Score"].round(2)
    subject_sheet = workbook.create_sheet("Performance By Subject")
    subject_sheet.cell(
        row=1, column=1,
        value="Pivot ready summary table. Values calculated by the build script from fact_assessment. Each student took one assessment in one subject, so subject rows compare different groups of students.",
    ).font = Font(italic=True, color="6B7280")
    write_frame(subject_sheet, by_subject, start_row=3)
    subject_sheet.freeze_panes = "A4"
    set_widths(subject_sheet, [12, 11, 15, 13, 15])

    # -------------------------------------------------------- Data Dictionary
    dictionary_rows = [
        ("Student Summary", "Student_ID", "Unique student identifier. Links every table.", "Text", "students.csv", "Unique inside the population."),
        ("Student Summary", "Grade", "Grade level, 5 to 8. Taken from the first fact table that records it for the student.", "Number", "attendance.csv, assessments.csv, discipline.csv", "Grade agrees across all three fact tables for every shared student."),
        ("Student Summary", "Program", "Program group: Core, MTSS Tier 1, MTSS Tier 2, or MTSS Tier 3.", "Text", "students.csv", "Casing standardized in the pipeline."),
        ("Student Summary", "Enrollment_Date", "Date the student enrolled.", "Date", "students.csv", "Runs from September to November 2025."),
        ("Student Summary", "Days_Enrolled", "Days the student was enrolled in the attendance period.", "Number", "attendance.csv", "Blank when the student has no attendance row."),
        ("Student Summary", "Days_Present", "Days the student attended.", "Number", "attendance.csv", "Never exceeds Days_Enrolled in the processed data."),
        ("Student Summary", "Days_Absent", "Days the student was absent.", "Number", "attendance.csv", "Present plus absent equals enrolled on every processed row."),
        ("Student Summary", "Attendance_Rate", "Days present as a percent of days enrolled, one decimal place.", "Number", "attendance.csv", "Basis for the bands and the chronic absence flag."),
        ("Student Summary", "Attendance_Band", "Project category from Attendance_Rate: At Risk, Needs Monitoring, On Track, or Strong.", "Text", "Derived in the pipeline", "Project categories, not official thresholds."),
        ("Student Summary", "Is_Chronic_Absence", "Yes when Attendance_Rate is below 90, the project definition.", "Text", "Derived in the pipeline", "Not an official school or district definition."),
        ("Student Summary", "Assessment", "Subject of the one assessment the student took: ELA, Math, or Science.", "Text", "assessments.csv", "One assessment per student."),
        ("Student Summary", "Assessment_Date", "Date of the assessment.", "Date", "assessments.csv", "Runs from January to April 2026."),
        ("Student Summary", "Score", "Assessment score from 43 to 100.", "Number", "assessments.csv", "Blank when the student has no assessment row."),
        ("Student Summary", "Benchmark", "Met when Score is 70 or above, Below when under 70.", "Text", "assessments.csv", "Rule verified on every processed row."),
        ("Student Summary", "Incidents", "Total discipline incidents recorded for the student.", "Number", "discipline.csv", "Blank when the student has no discipline row."),
        ("Student Summary", "Major_Incidents", "Major discipline incidents recorded for the student.", "Number", "discipline.csv", "A subset of Incidents."),
    ]
    dictionary = pd.DataFrame(
        dictionary_rows,
        columns=["Table", "Field", "Definition", "Data Type", "Source", "Notes"],
    )
    dictionary_sheet = workbook.create_sheet("Data Dictionary")
    write_frame(dictionary_sheet, dictionary)
    set_widths(dictionary_sheet, [17, 20, 62, 10, 34, 52])
    for row_index in range(2, len(dictionary) + 2):
        dictionary_sheet.cell(row=row_index, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        dictionary_sheet.cell(row=row_index, column=6).alignment = Alignment(wrap_text=True, vertical="top")

    # -------------------------------------------------------- Quality Summary
    quality_sheet = workbook.create_sheet("Quality Summary")
    quality_sheet.cell(
        row=1, column=1,
        value="Counts before and after the cleaning pipeline. Before counts come from sql/06_data_validation.sql run on the source files. After counts come from outputs/validation_results.csv.",
    ).font = Font(italic=True, color="6B7280")
    quality_rows = [
        ("Fact rows outside the analysis population, attendance", 11, 0),
        ("Fact rows outside the analysis population, assessments", 11, 0),
        ("Fact rows outside the analysis population, discipline", 10, 0),
        ("Program values with inconsistent casing", 1, 0),
        ("Benchmark values with inconsistent casing", 1, 0),
        ("Attendance rows where present plus absent differs from enrolled", 0, 0),
        ("Assessment rows breaking the benchmark rule", 0, 0),
    ]
    quality = pd.DataFrame(quality_rows, columns=["Check", "Before", "After"])
    write_frame(quality_sheet, quality, start_row=3)
    quality_sheet.freeze_panes = "A4"
    set_widths(quality_sheet, [64, 10, 10])
    note_row = len(quality) + 5
    passed = int((validation["result"] == "Pass").sum())
    quality_sheet.cell(
        row=note_row, column=1,
        value=f"The pipeline validation suite records {passed} of {len(validation)} checks passing on the processed tables. The full list sits in outputs/validation_results.csv.",
    )

    workbook.save(WORKBOOK_PATH)
    print(f"Wrote {WORKBOOK_PATH}")


if __name__ == "__main__":
    build_workbook()
